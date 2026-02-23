#!/usr/bin/env python3
"""
Convert all data files from proprietary formats to open source formats.
- DOCX → JSON or TXT
- XLSX → CSV

Maintains the original folder structure in the output directory.

Usage:
    python scripts/convert_to_open_formats.py

Output:
    converted_data/send_23_12_25/
"""

import json
import shutil
from pathlib import Path
from docx import Document
import pandas as pd


# Configuration - adjust these paths as needed
BASE_DIR = Path(__file__).parent.parent
SOURCE_DIR = BASE_DIR / "data_llm" / "send_23_12_25"
OUTPUT_DIR = BASE_DIR / "converted_data" / "send_23_12_25"


def clean_json_text(text: str) -> str:
    """Clean up JSON formatting issues from Word documents."""
    text = text.replace('\u00a0', ' ')
    text = text.replace('„', '"').replace('"', '"').replace('"', '"')
    text = text.replace(''', "'").replace(''', "'")
    text = text.replace('\t', '  ')
    text = text.replace('\u200b', '').replace('\ufeff', '')
    return text


def find_all_json_objects(text: str) -> list:
    """Find all top-level JSON objects in text using balanced brace matching."""
    text = clean_json_text(text)
    json_objects = []
    brace_count = 0
    start_idx = None

    for i, char in enumerate(text):
        if char == '{':
            if brace_count == 0:
                start_idx = i
            brace_count += 1
        elif char == '}':
            brace_count -= 1
            if brace_count == 0 and start_idx is not None:
                json_str = text[start_idx:i+1]
                try:
                    obj = json.loads(json_str)
                    json_objects.append(obj)
                except json.JSONDecodeError:
                    pass
                start_idx = None
    return json_objects


def extract_cases_from_object(obj: dict) -> list:
    """Extract case objects from a JSON object.

    Recognizes 4 schema formats:
    - v1.1: top-level case_meta + entitaeten
    - v1.0: wrapped in case_template
    - Diagnosen schema: top-level diagnosen + tumor_history + therapieplanung
    - English schema: top-level 'case' key with English field names (mRCC_case_schema)
    """
    cases = []
    if isinstance(obj, dict):
        if 'cases' in obj and isinstance(obj['cases'], list):
            cases.extend(obj['cases'])
        elif 'case_meta' in obj:
            cases.append(obj)
        elif 'case_template' in obj:
            cases.append(obj)
        elif 'diagnosen' in obj:
            # Diagnosen schema (v1.2-style) — cases with diagnosen instead of entitaeten
            cases.append(obj)
        elif 'case' in obj and isinstance(obj['case'], dict):
            # English schema (mRCC_case_schema) — case data wrapped in 'case' key
            cases.append(obj)
    return cases


def merge_cases_to_unified_structure(json_objects: list) -> dict:
    """Merge all JSON objects into a unified cases structure."""
    all_cases = []
    schema_version = "1.1"
    hinweis = None

    for obj in json_objects:
        if 'schema_version' in obj:
            schema_version = obj['schema_version']
        if 'hinweis_missing_values' in obj:
            hinweis = obj['hinweis_missing_values']
        cases = extract_cases_from_object(obj)
        all_cases.extend(cases)

    result = {"schema_version": schema_version}
    if hinweis:
        result["hinweis_missing_values"] = hinweis
    result["cases"] = all_cases
    return result


def convert_docx(docx_path: Path, output_path: Path) -> dict:
    """Convert DOCX to JSON or TXT."""
    result = {'source': str(docx_path.name), 'output': None, 'type': None, 'status': 'unknown'}

    try:
        doc = Document(str(docx_path))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        full_text = '\n'.join(paragraphs)

        if '{' not in full_text:
            # Plain text
            output_file = output_path.with_suffix('.txt')
            output_file.parent.mkdir(parents=True, exist_ok=True)
            with open(output_file, 'w', encoding='utf-8') as f:
                f.write(full_text)
            result['output'] = output_file.name
            result['type'] = 'txt'
            result['status'] = 'ok'
        else:
            # JSON content
            json_objects = find_all_json_objects(full_text)

            if json_objects:
                has_cases = any(
                    'case_meta' in obj or 'cases' in obj or 'case_template' in obj
                    or 'diagnosen' in obj
                    or ('case' in obj and isinstance(obj.get('case'), dict))
                    for obj in json_objects
                )

                if has_cases:
                    data = merge_cases_to_unified_structure(json_objects)
                else:
                    data = json_objects[0] if len(json_objects) == 1 else json_objects

                output_file = output_path.with_suffix('.json')
                output_file.parent.mkdir(parents=True, exist_ok=True)
                with open(output_file, 'w', encoding='utf-8') as f:
                    json.dump(data, f, ensure_ascii=False, indent=2)
                result['output'] = output_file.name
                result['type'] = 'json'
                result['status'] = 'ok'

                if has_cases and 'cases' in data:
                    result['case_count'] = len(data['cases'])
            else:
                result['status'] = 'no_json'

    except Exception as e:
        result['status'] = f'error: {e}'

    return result


def convert_xlsx(xlsx_path: Path, output_path: Path) -> dict:
    """Convert XLSX to CSV."""
    result = {'source': str(xlsx_path.name), 'output': None, 'type': 'csv', 'status': 'unknown'}

    try:
        df = pd.read_excel(xlsx_path)
        output_file = output_path.with_suffix('.csv')
        output_file.parent.mkdir(parents=True, exist_ok=True)
        df.to_csv(output_file, index=False, encoding='utf-8')

        result['output'] = output_file.name
        result['status'] = 'ok'
        result['rows'] = len(df)

    except Exception as e:
        result['status'] = f'error: {e}'

    return result


def main():
    # Clean output directory
    if OUTPUT_DIR.exists():
        shutil.rmtree(OUTPUT_DIR)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("=" * 70)
    print("CONVERTING ALL FILES TO OPEN FORMATS")
    print("=" * 70)
    print(f"\nSource: {SOURCE_DIR}")
    print(f"Output: {OUTPUT_DIR}")

    # Find all files
    docx_files = list(SOURCE_DIR.rglob("*.docx"))
    xlsx_files = list(SOURCE_DIR.rglob("*.xlsx"))

    print(f"\nFound {len(docx_files)} DOCX files and {len(xlsx_files)} XLSX files")

    results = []

    # Convert DOCX files
    print("\n" + "-" * 70)
    print("CONVERTING DOCX FILES")
    print("-" * 70 + "\n")

    for docx_file in sorted(docx_files):
        rel_path = docx_file.relative_to(SOURCE_DIR)
        output_path = OUTPUT_DIR / rel_path.with_suffix('')

        print(f"Converting: {rel_path}")
        result = convert_docx(docx_file, output_path)
        results.append(result)

        if result['status'] == 'ok':
            extra = f" ({result.get('case_count', 0)} cases)" if result.get('case_count') else ""
            print(f"  ✓ → {result['output']}{extra}")
        else:
            print(f"  ✗ {result['status']}")

    # Convert XLSX files
    print("\n" + "-" * 70)
    print("CONVERTING XLSX FILES")
    print("-" * 70 + "\n")

    for xlsx_file in sorted(xlsx_files):
        rel_path = xlsx_file.relative_to(SOURCE_DIR)
        output_path = OUTPUT_DIR / rel_path.with_suffix('')

        print(f"Converting: {rel_path}")
        result = convert_xlsx(xlsx_file, output_path)
        results.append(result)

        if result['status'] == 'ok':
            print(f"  ✓ → {result['output']} ({result['rows']} rows)")
        else:
            print(f"  ✗ {result['status']}")

    # Summary
    print("\n" + "=" * 70)
    print("CONVERSION COMPLETE")
    print("=" * 70)

    ok_count = sum(1 for r in results if r['status'] == 'ok')
    print(f"\n✓ {ok_count}/{len(results)} files converted successfully")


if __name__ == "__main__":
    main()
